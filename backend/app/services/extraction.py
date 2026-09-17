"""Pull plain text out of an uploaded document.

Supports PDF, DOCX, and plain text. The original app only ever captured the
filename, so attachments never reached the model at all; this is what makes
"attach your existing rubric" mean something.
"""

import io
import os
from typing import Tuple

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md", ".text", ".csv")


class ExtractionError(ValueError):
    """Raised when a file cannot be read as text."""


def extract_text(filename: str, data: bytes, max_chars: int) -> Tuple[str, bool]:
    """Return (text, truncated) for an uploaded file.

    Raises ExtractionError with a teacher-readable message on unsupported or
    unreadable files.
    """
    extension = os.path.splitext(filename or "")[1].lower()

    if extension == ".pdf":
        text = _from_pdf(data)
    elif extension == ".docx":
        text = _from_docx(data)
    elif extension in (".txt", ".md", ".text", ".csv"):
        text = _from_plain(data)
    elif extension == ".doc":
        raise ExtractionError(
            "Legacy .doc files can't be read. Save the file as .docx or PDF "
            "and attach it again."
        )
    else:
        raise ExtractionError(
            "Unsupported file type '%s'. Attach a PDF, DOCX, or text file."
            % (extension or "unknown")
        )

    text = _normalize(text)
    if not text:
        raise ExtractionError(
            "No text could be read from this file. If it is a scanned document, "
            "the pages are images and would need OCR."
        )

    if len(text) > max_chars:
        return text[:max_chars].rstrip() + "\n\n[document truncated]", True
    return text, False


def _from_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover - dependency is declared
        raise ExtractionError("PDF support is not installed (pip install pypdf).")

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("Could not open the PDF: %s" % exc)

    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception:
            raise ExtractionError(
                "This PDF is password protected. Remove the protection and "
                "attach it again."
            )

    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return "\n\n".join(pages)


def _from_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError:  # pragma: no cover - dependency is declared
        raise ExtractionError("DOCX support is not installed (pip install python-docx).")

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("Could not open the Word document: %s" % exc)

    blocks = [p.text for p in document.paragraphs]

    # Existing rubrics are almost always tables, so read those too.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))

    return "\n".join(blocks)


def _from_plain(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Could not decode this file as text.")


def _normalize(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").split("\n")]

    # Collapse runs of blank lines to at most one.
    out, blank = [], False
    for line in lines:
        if line.strip():
            out.append(line)
            blank = False
        elif not blank:
            out.append("")
            blank = True
    return "\n".join(out).strip()
