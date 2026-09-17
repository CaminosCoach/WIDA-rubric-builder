"""Build a genuine .docx from a generated rubric.

The original app produced an HTML string, renamed it .doc, and let Word guess —
which is why Word warned about the file on open. This writes a real Office
Open XML document with real tables.
"""

import io
from typing import Any, Dict, List

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INK = RGBColor(0x18, 0x1F, 0x2E)
AMBER = RGBColor(0x93, 0x61, 0x1F)
SOFT = RGBColor(0x5A, 0x5E, 0x4F)
SAGE = RGBColor(0x3F, 0x5C, 0x4C)

# Row fills, mirroring the level dots in the web view.
LEVEL_FILLS = {
    6: "181F2E", 5: "93611F", 4: "C68A3D",
    3: "DCB87C", 2: "E9D7AE", 1: "F1E6C8",
}


def build_docx(rubric: Dict[str, Any], meta: Dict[str, Any]) -> bytes:
    document = Document()
    _setup_page(document)

    _header(document, rubric, meta)
    _content_objectives(document, rubric.get("contentObjectives") or {})
    _language_objectives(document, rubric.get("languageObjectives") or {})
    _rubric_table(document, "03", "Interpretive / Process rubric",
                  rubric.get("interpretiveRubric") or {})
    _rubric_table(document, "04", "Expressive / Product rubric",
                  rubric.get("expressiveRubric") or {})
    _content_evidence(document, rubric.get("contentEvidence") or {})
    _scaffolds(document, rubric.get("scaffolds") or {})
    _assumptions(document, rubric.get("assumptions") or {})

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------
# page + text helpers
# --------------------------------------------------------------------------

def _setup_page(document: Document) -> None:
    section = document.sections[0]
    section.left_margin = section.right_margin = Inches(0.7)
    section.top_margin = section.bottom_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)


def _para(document, text="", size=10.5, bold=False, color=None,
          space_after=6, space_before=0, italic=False):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.space_before = Pt(space_before)
    if text:
        run = paragraph.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        if color is not None:
            run.font.color.rgb = color
    return paragraph


def _section_heading(document, number: str, title: str, note: str = "") -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(18)
    paragraph.paragraph_format.space_after = Pt(2)

    num_run = paragraph.add_run(number + "  ")
    num_run.font.size = Pt(9)
    num_run.bold = True
    num_run.font.color.rgb = AMBER

    title_run = paragraph.add_run(title)
    title_run.font.size = Pt(15)
    title_run.bold = True
    title_run.font.color.rgb = INK

    if note:
        _para(document, note, size=9.5, color=SOFT, space_after=10)


def _shade(cell, hex_fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _cell_text(cell, text: str, size=9, bold=False, color=None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text or "")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _bordered_table(document, rows: int, cols: int):
    table = document.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    return table


# --------------------------------------------------------------------------
# sections
# --------------------------------------------------------------------------

def _header(document, rubric: Dict[str, Any], meta: Dict[str, Any]) -> None:
    _para(document, meta.get("kicker", ""), size=9.5, bold=True,
          color=AMBER, space_after=4)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run(rubric.get("title") or "WIDA-Aligned Rubric")
    run.font.size = Pt(20)
    run.bold = True
    run.font.color.rgb = INK

    question = document.add_paragraph()
    question.paragraph_format.space_after = Pt(12)
    label = question.add_run("Focusing question: ")
    label.bold = True
    label.font.size = Pt(10.5)
    body = question.add_run(meta.get("focusingQuestion", "") or "—")
    body.font.size = Pt(10.5)
    body.font.color.rgb = SOFT

    pairs: List = [
        ("Anchor text", meta.get("anchorText")),
        ("Unit", meta.get("unit")),
        ("Standards", meta.get("standards")),
        ("English Learners", meta.get("englishLearners")),
    ]
    table = _bordered_table(document, rows=2, cols=len(pairs))
    for index, (label_text, value) in enumerate(pairs):
        _cell_text(table.cell(0, index), label_text, size=8, bold=True, color=SOFT)
        _cell_text(table.cell(1, index), value or "—", size=9.5, bold=True)


def _content_objectives(document, data: Dict[str, Any]) -> None:
    _section_heading(document, "01", "Content objectives", data.get("note", ""))

    items = data.get("items") or []
    if not items:
        return
    table = _bordered_table(document, rows=len(items), cols=2)
    table.columns[0].width = Inches(2.6)
    table.columns[1].width = Inches(4.5)
    for row_index, item in enumerate(items):
        _cell_text(table.cell(row_index, 0), item.get("label", ""), size=9.5, bold=True)
        _cell_text(table.cell(row_index, 1), item.get("detail", ""), size=9.5, color=SOFT)


def _language_objectives(document, data: Dict[str, Any]) -> None:
    _section_heading(document, "02",
                     "Interpretive & expressive language objectives",
                     data.get("note", ""))

    columns = [data.get("interpretive") or {}, data.get("expressive") or {}]
    table = _bordered_table(document, rows=1, cols=2)
    for index, column in enumerate(columns):
        cell = table.cell(0, index)
        cell.text = ""

        heading = cell.paragraphs[0]
        heading_run = heading.add_run(column.get("heading", ""))
        heading_run.bold = True
        heading_run.font.size = Pt(9.5)
        heading_run.font.color.rgb = AMBER

        for dimension in ("discourse", "sentence", "word"):
            label_para = cell.add_paragraph()
            label_para.paragraph_format.space_after = Pt(0)
            label_para.paragraph_format.space_before = Pt(6)
            label_run = label_para.add_run(
                "Word/Phrase" if dimension == "word" else dimension.capitalize())
            label_run.bold = True
            label_run.font.size = Pt(8.5)
            label_run.font.color.rgb = SOFT

            text_para = cell.add_paragraph()
            text_para.paragraph_format.space_after = Pt(0)
            text_run = text_para.add_run(column.get(dimension, ""))
            text_run.font.size = Pt(9.5)


def _rubric_table(document, number: str, title: str, data: Dict[str, Any]) -> None:
    _section_heading(document, number, title, data.get("note", ""))

    rows = data.get("rows") or []
    if not rows:
        return

    table = _bordered_table(document, rows=len(rows) + 1, cols=4)
    headers = [
        "Level",
        "Discourse — Organization & cohesion",
        "Sentence — Grammatical complexity",
        "Word/Phrase — Precision",
    ]
    for index, header in enumerate(headers):
        cell = table.cell(0, index)
        _shade(cell, "181F2E")
        _cell_text(cell, header, size=8.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    table.columns[0].width = Inches(0.95)
    for index in range(1, 4):
        table.columns[index].width = Inches(2.05)

    for row_index, row in enumerate(rows, start=1):
        level = int(row.get("level", 0) or 0)

        level_cell = table.cell(row_index, 0)
        level_cell.text = ""
        level_para = level_cell.paragraphs[0]
        level_para.paragraph_format.space_after = Pt(0)
        level_run = level_para.add_run(str(level))
        level_run.bold = True
        level_run.font.size = Pt(11)
        name_para = level_cell.add_paragraph()
        name_para.paragraph_format.space_after = Pt(0)
        name_run = name_para.add_run(row.get("name", ""))
        name_run.font.size = Pt(8)
        name_run.font.color.rgb = SOFT
        if level in LEVEL_FILLS:
            _shade(level_cell, LEVEL_FILLS[level])
            if level >= 5:
                level_run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                name_run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        _cell_text(table.cell(row_index, 1), row.get("discourse", ""))
        _cell_text(table.cell(row_index, 2), row.get("sentence", ""))

        word_cell = table.cell(row_index, 3)
        _cell_text(word_cell, row.get("word", ""))
        examples = [e for e in (row.get("examples") or []) if e]
        if examples:
            example_para = word_cell.add_paragraph()
            example_para.paragraph_format.space_before = Pt(2)
            example_para.paragraph_format.space_after = Pt(0)
            example_run = example_para.add_run("e.g., " + ", ".join(examples))
            example_run.font.size = Pt(8)
            example_run.font.color.rgb = SAGE


def _content_evidence(document, data: Dict[str, Any]) -> None:
    _section_heading(document, "05", "Content evidence checklist", data.get("note", ""))
    for item in data.get("items") or []:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.left_indent = Inches(0.12)
        box = paragraph.add_run("☐  ")
        box.font.size = Pt(12)
        box.font.color.rgb = SAGE
        text = paragraph.add_run(item)
        text.font.size = Pt(9.5)


def _scaffolds(document, data: Dict[str, Any]) -> None:
    _section_heading(document, "06", "Scaffolds & supports", data.get("note", ""))

    bands = data.get("bands") or []
    if not bands:
        return

    table = _bordered_table(document, rows=1, cols=len(bands))
    for index, band in enumerate(bands):
        cell = table.cell(0, index)
        cell.text = ""

        heading = cell.paragraphs[0]
        heading.paragraph_format.space_after = Pt(4)
        heading_run = heading.add_run(band.get("label", ""))
        heading_run.bold = True
        heading_run.font.size = Pt(9.5)
        heading_run.font.color.rgb = AMBER

        for item in band.get("items") or []:
            item_para = cell.add_paragraph()
            item_para.paragraph_format.space_after = Pt(1)
            item_para.paragraph_format.left_indent = Inches(0.1)
            item_run = item_para.add_run("•  " + item)
            item_run.font.size = Pt(9)
            item_run.font.color.rgb = SOFT


def _assumptions(document, data: Dict[str, Any]) -> None:
    _section_heading(document, "07", "Assumptions & implementation notes")
    for item in data.get("items") or []:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        tag = paragraph.add_run("[%s]  " % item.get("tag", "note"))
        tag.font.size = Pt(8.5)
        tag.bold = True
        tag.font.color.rgb = SAGE
        text = paragraph.add_run(item.get("note", ""))
        text.font.size = Pt(9.5)
        text.font.color.rgb = SOFT


def build_filename(grade_level: str) -> str:
    safe = "".join(c if c.isalnum() else "_" for c in (grade_level or "Rubric"))
    safe = "_".join(part for part in safe.split("_") if part) or "Rubric"
    return "WIDA_Rubric_%s.docx" % safe
